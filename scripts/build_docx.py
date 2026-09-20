"""Convierte report/report.md en report/report.docx (markdown-it + python-docx).

Genera un Word nativo: títulos con estilos de encabezado (aparecen en el panel de navegación),
tablas reales, figuras incrustadas, bloques de código en monoespaciado, número de página y
español como idioma del documento.
"""
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from markdown_it import MarkdownIt

from src.config import ROOT

REPORT_DIR = ROOT / "report"
SOURCE = REPORT_DIR / "report.md"
OUTPUT = REPORT_DIR / "report.docx"

BODY_FONT, HEADING_FONT, CODE_FONT = "Georgia", "Calibri", "Consolas"
CODE_FILL, HEADER_FILL = "F4F3F0", "E9E8E4"
IMAGE_WIDTH = Cm(16)


def set_font(style_or_run, name: str) -> None:
    style_or_run.font.name = name
    rpr = style_or_run.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if fonts.get(qn(attr)) is not None:
            del fonts.attrib[qn(attr)]  # si quedan, Word ignora el tipo de letra elegido
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)


def set_language(style, code: str = "es-ES") -> None:
    rpr = style.element.get_or_add_rPr()
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), code)
    lang.set(qn("w:eastAsia"), code)


def shade(element_pr, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    element_pr.append(shd)


def add_page_number(section) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
        else:
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = text
        run._r.append(element)
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def make_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.0)
    section.top_margin, section.bottom_margin = Cm(2.0), Cm(2.2)
    add_page_number(section)

    normal = doc.styles["Normal"]
    set_font(normal, BODY_FONT)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    set_language(normal)
    for name, size in (("Title", 21), ("Heading 1", 14), ("Heading 2", 12)):
        style = doc.styles[name]
        set_font(style, HEADING_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(16)
    doc.styles["Heading 1"].paragraph_format.keep_with_next = True
    return doc


def add_inline(paragraph, inline, size=None, bold=False, italic=False) -> None:
    """Vuelca los hijos de un token `inline` como runs con su formato."""
    state = {"bold": bold, "italic": italic, "code": False}
    link_href = None
    for child in inline.children or []:
        kind = child.type
        if kind == "strong_open":
            state["bold"] = True
        elif kind == "strong_close":
            state["bold"] = bold
        elif kind == "em_open":
            state["italic"] = True
        elif kind == "em_close":
            state["italic"] = italic
        elif kind == "link_open":
            link_href = child.attrGet("href")
        elif kind == "link_close":
            if link_href:
                run = paragraph.add_run(f" ({link_href})")
                run.font.size = Pt(size) if size else None
            link_href = None
        elif kind in ("text", "code_inline"):
            run = paragraph.add_run(child.content)
            run.bold, run.italic = state["bold"], state["italic"]
            if size:
                run.font.size = Pt(size)
            if kind == "code_inline":
                set_font(run, CODE_FONT)
                run.font.size = Pt((size or 10.5) - 1)
        elif kind == "softbreak":
            paragraph.add_run(" ")
        elif kind == "hardbreak":
            paragraph.add_run().add_break()
        elif kind == "image":
            pass  # las imágenes se manejan a nivel de párrafo
        # html_inline (comentarios) se ignora


def image_of(inline):
    for child in inline.children or []:
        if child.type == "image":
            return child.attrGet("src")
    return None


def restart_numbering(doc, style_name: str = "List Number") -> int:
    """Crea una numeración nueva (empieza en 1) sobre la definición del estilo, para que cada lista numerada reinicie."""
    numbering = doc.part.numbering_part.element
    style_num = doc.styles[style_name].element.pPr.numPr.numId.val
    abstract = next(n for n in numbering.findall(qn("w:num")) if int(n.get(qn("w:numId"))) == style_num).find(qn("w:abstractNumId")).get(qn("w:val"))
    new_id = max(int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(new_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.extend([ref, override])
    numbering.append(num)
    return new_id


def use_numbering(paragraph, num_id: int) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    ident = OxmlElement("w:numId")
    ident.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, ident])
    ppr.append(num_pr)


class Converter:
    def __init__(self, base_dir: Path):
        self.doc = make_document()
        self.base_dir = base_dir
        self.lists: list[dict] = []

    def convert(self, markdown_text: str) -> Document:
        tokens = MarkdownIt("commonmark", {"html": True}).enable("table").parse(markdown_text)
        i = 0
        while i < len(tokens):
            i = self.handle(tokens, i)
        return self.doc

    def handle(self, tokens, i) -> int:
        token = tokens[i]
        kind = token.type
        if kind == "heading_open":
            level, inline = token.tag, tokens[i + 1]
            style = "Title" if level == "h1" else "Heading 1" if level == "h2" else "Heading 2"
            paragraph = self.doc.add_paragraph(style=style)
            add_inline(paragraph, inline)
            # También a nivel de run: algunos visores (Vista Previa, Pages, Google Docs) ignoran el formato del estilo.
            for run in paragraph.runs:
                run.bold = True
                set_font(run, HEADING_FONT)
                run.font.size = Pt({"Title": 21, "Heading 1": 14}.get(style, 12))
            return i + 3
        if kind == "paragraph_open":
            return self.paragraph(tokens, i)
        if kind in ("bullet_list_open", "ordered_list_open"):
            ordered = kind == "ordered_list_open"
            self.lists.append({"ordered": ordered, "n": 0, "num_id": restart_numbering(self.doc) if ordered else None})
        elif kind in ("bullet_list_close", "ordered_list_close"):
            self.lists.pop()
        elif kind == "list_item_open":
            self.lists[-1]["n"] += 1
        elif kind == "fence":
            self.code_block(token.content)
        elif kind == "table_open":
            return self.table(tokens, i)
        return i + 1  # html_block (comentarios), hr, cierres: sin efecto

    def paragraph(self, tokens, i) -> int:
        inline = tokens[i + 1]
        src = image_of(inline)
        if src and not self.lists:
            path = (self.base_dir / src).resolve()
            self.doc.add_picture(str(path), width=IMAGE_WIDTH)
            self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            self.doc.paragraphs[-1].paragraph_format.keep_with_next = True
            return i + 3
        caption = re.match(r"^\*Figura \d+\.", inline.content or "")
        if self.lists:
            ctx = self.lists[-1]
            paragraph = self.doc.add_paragraph(style="List Number" if ctx["ordered"] else "List Bullet")
            if ctx["ordered"]:
                use_numbering(paragraph, ctx["num_id"])
            paragraph.paragraph_format.space_after = Pt(3)
        else:
            paragraph = self.doc.add_paragraph()
        if caption:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_inline(paragraph, inline, size=9)
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
        else:
            add_inline(paragraph, inline)
            if not self.lists:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        return i + 3

    def code_block(self, content: str) -> None:
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.3)
        paragraph.paragraph_format.space_after = Pt(8)
        shade(paragraph._p.get_or_add_pPr(), CODE_FILL)
        lines = content.rstrip("\n").split("\n")
        for n, line in enumerate(lines):
            run = paragraph.add_run(line)
            set_font(run, CODE_FONT)
            run.font.size = Pt(8.5)
            if n < len(lines) - 1:
                run.add_break()

    def table(self, tokens, i) -> int:
        rows, header, j = [], set(), i + 1
        while tokens[j].type != "table_close":
            if tokens[j].type == "tr_open":
                cells = []
            elif tokens[j].type in ("th_open", "td_open"):
                if tokens[j].type == "th_open":
                    header.add(len(rows))
                cells.append(tokens[j + 1])
            elif tokens[j].type == "tr_close":
                rows.append(cells)
            j += 1
        table = self.doc.add_table(rows=len(rows), cols=len(rows[0]))
        table.style = "Table Grid"
        table.autofit = True
        for r, cells in enumerate(rows):
            for c, inline in enumerate(cells):
                cell = table.cell(r, c)
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                add_inline(paragraph, inline, size=8, bold=r in header)
                if r in header:
                    shade(cell._tc.get_or_add_tcPr(), HEADER_FILL)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return j + 1


def student_name(markdown_text: str) -> str:
    match = re.search(r"\*\*Estudiante:\*\*\s*(.+?)\s*$", markdown_text, re.M)
    return match.group(1) if match else ""


def main() -> None:
    if not SOURCE.exists():
        raise SystemExit("Falta report/report.md: correr scripts/build_report.py primero.")
    text = SOURCE.read_text(encoding="utf-8")
    doc = Converter(REPORT_DIR).convert(text)
    props = doc.core_properties
    props.title = "Taller 01 — Foundation Models"
    props.author = student_name(text)
    props.language = "es-ES"
    doc.save(OUTPUT)
    print(f"{OUTPUT.relative_to(ROOT)} generado ({OUTPUT.stat().st_size / 1024:.0f} KB)")


if __name__ == "__main__":
    main()
